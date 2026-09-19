import json
import os
import random
import re
import sys
import threading
import time
import urllib.parse
from datetime import datetime
from tkinter import filedialog, messagebox

import customtkinter as ctk
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import sync_playwright

try:
    from torob_chart import TorobChartModule
except ImportError:
    TorobChartModule = None


def get_resource_path(relative_path):
    if hasattr(sys, "_MEIPASS"):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)


ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")


class TorobScraperApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("سامانه استخراج داده ترب")
        self.geometry("720x800")
        self.resizable(False, False)

        try:
            icon_path = get_resource_path("icon.ico")
            if os.path.exists(icon_path):
                self.iconbitmap(icon_path)
        except Exception:  # noqa: BLE001, S110
            pass

        self.is_running = False
        self.stop_requested = False
        self.selected_file_path = ""
        self.mode_var = ctk.StringVar(value="استخراج از فایل اکسل")

        self.create_widgets()

    def create_widgets(self):
        self.lbl_title = ctk.CTkLabel(
            self,
            text="📊 استخراج‌گر هوشمند اطلاعات ترب",
            font=ctk.CTkFont(family="Vazirmatn", size=22, weight="bold"),
        )
        self.lbl_title.pack(pady=(15, 10))

        # فریم انتخاب حالت برنامه
        self.frame_mode = ctk.CTkFrame(self)
        self.frame_mode.pack(padx=20, pady=(0, 10), fill="x")

        self.seg_button = ctk.CTkSegmentedButton(
            self.frame_mode,
            values=["استخراج از فایل اکسل", "جستجوی تکی"],
            command=self.toggle_mode,
            variable=self.mode_var,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.seg_button.pack(padx=10, pady=10, fill="x")

        # فریم‌های حالت اکسل
        self.frame_file = ctk.CTkFrame(self)
        self.btn_select_file = ctk.CTkButton(
            self.frame_file,
            text="انتخاب فایل اکسل ورودی",
            command=self.select_file,
            width=180,
        )
        self.btn_select_file.pack(side="right", padx=10, pady=10)

        self.lbl_file_path = ctk.CTkLabel(
            self.frame_file, text="هیچ فایلی انتخاب نشده است", text_color="gray"
        )
        self.lbl_file_path.pack(side="left", padx=10, pady=10)

        self.frame_rows = ctk.CTkFrame(self)
        self.lbl_end = ctk.CTkLabel(self.frame_rows, text="تا سطر:")
        self.lbl_end.pack(side="right", padx=(2, 8), pady=10)

        self.entry_end_row = ctk.CTkEntry(self.frame_rows, width=60)
        self.entry_end_row.insert(0, "10")
        self.entry_end_row.pack(side="right", padx=(0, 10), pady=10)

        self.lbl_start = ctk.CTkLabel(self.frame_rows, text="از سطر:")
        self.lbl_start.pack(side="right", padx=(2, 5), pady=10)

        self.entry_start_row = ctk.CTkEntry(self.frame_rows, width=60)
        self.entry_start_row.insert(0, "1")
        self.entry_start_row.pack(side="right", padx=(0, 10), pady=10)

        # فریم حالت جستجوی تکی
        self.frame_single = ctk.CTkFrame(self)
        self.lbl_single = ctk.CTkLabel(self.frame_single, text="نام محصول در ترب:")
        self.lbl_single.pack(side="right", padx=(10, 5), pady=15)

        self.entry_single = ctk.CTkEntry(self.frame_single, width=400)
        self.entry_single.pack(side="right", padx=(5, 10), pady=15)

        # فریم تنظیمات عمومی و اکشن‌ها
        self.frame_actions = ctk.CTkFrame(self)
        self.frame_actions.pack(padx=20, pady=8, fill="x")

        self.switch_headless = ctk.CTkSwitch(self.frame_actions, text="مرورگر مخفی")
        self.switch_headless.pack(side="right", padx=15, pady=10)
        self.switch_headless.deselect()

        self.btn_stop = ctk.CTkButton(
            self.frame_actions,
            text="🚨 توقف و ذخیره",
            fg_color="#dc3545",
            hover_color="#c82333",
            command=self.trigger_emergency_stop,
            height=40,
            width=140,
            state="disabled",
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.btn_stop.pack(side="left", padx=10, pady=10)

        self.btn_start = ctk.CTkButton(
            self.frame_actions,
            text="▶ شروع استخراج",
            fg_color="#28a745",
            hover_color="#218838",
            command=self.start_scraping_thread,
            height=40,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.btn_start.pack(side="left", padx=10, pady=10, expand=True, fill="x")

        self.txt_log = ctk.CTkTextbox(self, width=680, height=350)
        self.txt_log.pack(padx=20, pady=(8, 15))

        self.toggle_mode("استخراج از فایل اکسل")
        self.log("💡 آماده‌باش: حالت اجرای برنامه را انتخاب کرده و «شروع» را بزنید.")

    def toggle_mode(self, selected_mode):
        if selected_mode == "استخراج از فایل اکسل":
            self.frame_single.pack_forget()
            self.frame_file.pack(padx=20, pady=8, fill="x", after=self.frame_mode)
            self.frame_rows.pack(padx=20, pady=8, fill="x", after=self.frame_file)
        else:
            self.frame_file.pack_forget()
            self.frame_rows.pack_forget()
            self.frame_single.pack(padx=20, pady=8, fill="x", after=self.frame_mode)

    def log(self, message):
        self.txt_log.insert("end", f"{message}\n")
        self.txt_log.see("end")

    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="انتخاب فایل اکسل محصولات",
            filetypes=[("Excel Files", "*.xlsx *.xls")],
        )
        if file_path:
            self.selected_file_path = file_path
            filename = os.path.basename(file_path)
            self.lbl_file_path.configure(text=filename, text_color="white")
            self.log(f"📁 فایل ورودی انتخاب شد: {filename}")

    def trigger_emergency_stop(self):
        if self.is_running:
            self.stop_requested = True
            self.log("\n🚨 درخواست توقف دریافت شد... ذخیره‌سازی انجام می‌شود.")

    def handle_captcha_if_present(self, page):
        try:
            title = page.title().lower()
            current_url = page.url.lower()

            is_blocked = (
                "captcha" in current_url
                or "challenge" in current_url
                or "403" in title
                or "forbidden" in title
                or "تایید هویت" in title
                or "دسترسی غیرمجاز" in title
            )

            if is_blocked:
                self.log("\n⚠️ [سد امنیتی ترب] کپچای واقعی ظاهر شد!")
                self.log("👉 لطفاً کپچا را در مرورگر حل کنید...")

                messagebox.showwarning(
                    "تایید هویت",
                    "لطفاً کپچای ظاهرشده در مرورگر را حل کنید تا استخراج ادامه یابد.",
                )

                while not self.stop_requested:
                    time.sleep(1.5)
                    cur_title = page.title().lower()
                    cur_url = page.url.lower()
                    still_blocked = (
                        "captcha" in cur_url
                        or "403" in cur_title
                        or "forbidden" in cur_title
                        or "تایید هویت" in cur_title
                    )
                    if not still_blocked:
                        self.log("✅ کپچا حل شد! ادامه استخراج...\n")
                        time.sleep(1)
                        break
        except Exception:  # noqa: BLE001, S110
            pass

    def safe_fetch_json(self, page, url):
        js_fetch = """async (target_url) => {
                try {
                    const res = await fetch(target_url, {
                        headers: {
                            'accept': 'application/json, text/plain, */*',
                            'accept-language': 'fa-IR,fa;q=0.9,en;q=0.8'
                        }
                    });
                    if (res.status === 403 || res.status === 429) {
                        return { status: res.status, data: null };
                    }
                    if (!res.ok) return { status: res.status, data: null };
                    const data = await res.json();
                    return { status: 200, data: data };
                } catch(e) {
                    return { status: 500, data: null };
                }
            }"""

        if "torob.com" not in page.url or "api.torob.com" in page.url:
            page.goto(
                "https://torob.com/", wait_until="domcontentloaded", timeout=30000
            )

        res = page.evaluate(js_fetch, url)

        if res and res.get("status") in [403, 429]:
            page.goto(
                "https://torob.com/", wait_until="domcontentloaded", timeout=30000
            )
            self.handle_captcha_if_present(page)
            res = page.evaluate(js_fetch, url)

        return res.get("data") if res else None

    def price_data_cleaner(self, raw_price):
        if not raw_price or raw_price == "-":
            return None, "نامشخص"
        price_number_str = raw_price.replace(",", "").replace("،", "").replace("٫", "")
        status = "ناموجود" if "ناموجود" in raw_price else "موجود"
        price_number = re.findall(r"\d+", price_number_str)
        if price_number:
            return int(price_number[0]), status
        return None, status

    def cpt_data_cleaner(self, raw_cpt):
        """
        this function cleans cpt data(cpt = change price time)
        """
        if not raw_cpt or raw_cpt == "-":
            return "", ""

        if "دیروز" in raw_cpt:
            return 1, "روز"
        elif "پریروز" in raw_cpt:
            return 2, "روز"
        elif "لحظاتی" in raw_cpt or "دقایقی" in raw_cpt:
            return 0, "دقیقه"

        matches = re.findall(r"(\d+)\s*(سال|ماه|هفته|روز|ساعت|دقیقه|ثانیه)", raw_cpt)
        if not matches:
            return "", ""

        if len(matches) == 1:
            time_value, time_unit = matches[0]
            return int(time_value), time_unit
        elif len(matches) == 2:
            unit_data = {
                "سال": 365 * 24 * 60,
                "ماه": 30 * 24 * 60,
                "هفته": 7 * 24 * 60,
                "روز": 24 * 60,
                "ساعت": 60,
                "دقیقه": 1,
                "ثانیه": 1 / 60,
            }
            time_value1, time_unit1 = matches[0]
            time_value2, time_unit2 = matches[1]
            time_value1, time_value2 = int(time_value1), int(time_value2)

            total_minute = time_value1 * unit_data.get(
                time_unit1, 0
            ) + time_value2 * unit_data.get(time_unit2, 0)
            minute2unit2 = total_minute / unit_data.get(time_unit2, 1)
            return minute2unit2, time_unit2

    def wait_for_unblock(self):
        self.log("⏳ صبر یک ساعته به دلیل احتمال مسدودیت ترب...")
        elapsed = 0
        while elapsed < 3600 and not self.stop_requested:
            time.sleep(5)
            elapsed += 5

    def start_scraping_thread(self):
        mode = self.mode_var.get()
        search_query = None
        start_row = 1
        end_row = 1
        max_results = 10

        if mode == "استخراج از فایل اکسل":
            if not self.selected_file_path:
                messagebox.showwarning("خطا", "لطفاً ابتدا فایل اکسل را انتخاب کنید.")
                return
            try:
                start_row = int(self.entry_start_row.get())
                end_row = int(self.entry_end_row.get())
            except ValueError:
                messagebox.showerror("خطا", "لطفاً مقادیر معتبر وارد کنید.")
                return
            max_results = 10
        else:
            search_query = self.entry_single.get().strip()
            if not search_query:
                messagebox.showwarning("خطا", "لطفاً نام محصول را وارد کنید.")
                return
            max_results = 20

        self.is_running = True
        self.stop_requested = False
        self.btn_start.configure(state="disabled")
        self.btn_stop.configure(state="normal")
        self.seg_button.configure(state="disabled")
        self.btn_select_file.configure(state="disabled")

        is_headless = bool(self.switch_headless.get())

        threading.Thread(
            target=self.run_scraper,
            args=(mode, search_query, start_row, end_row, max_results, is_headless),
            daemon=True,
        ).start()

    def run_scraper(
        self, mode, search_query, start_row, end_row, max_results, is_headless
    ):
        try:
            now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")  # noqa: DTZ005

            if mode == "استخراج از فایل اکسل":
                input_df = pd.read_excel(self.selected_file_path)
                total_rows = len(input_df)
                sliced_df = input_df.iloc[start_row - 1 : end_row]
                base_filename = f"Torob_Report_Rows_{start_row}_to_{end_row}_{now}"
                self.log(
                    f"\n🚀 شروع استخراج سطرهای {start_row} تا {min(end_row, total_rows)}..."
                )
            else:
                sliced_df = pd.DataFrame([{"کد": "1", "نام": search_query}])
                safe_name = "".join(
                    x for x in search_query if x.isalnum() or x in " _-"
                )[:30]
                base_filename = f"Torob_Single_{safe_name}_{now}"
                self.log(f"\n🚀 شروع جستجوی تکی (۲۰ نتیجه) برای: {search_query}...")

            # یافتن پوشه اصلی که برنامه (سورس یا exe) در آن در حال اجراست
            if getattr(sys, "frozen", False):
                script_dir = os.path.dirname(sys.executable)
            else:
                script_dir = os.path.dirname(os.path.abspath(__file__))

            # تعیین مسیر قطعی و مطلق برای فایل‌های خروجی
            excel_filename = os.path.join(script_dir, f"{base_filename}.xlsx")
            json_filename = os.path.join(script_dir, f"{base_filename}.json")

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "گزارش ترب"

            if mode == "استخراج از فایل اکسل":
                input_col = input_df.columns.tolist()
            else:
                input_col = ["کد", "نام"]

            headers = input_col + [
                "نام در ترب",
                "فروشگاه",
                "امتیاز و سابقه",
                "ضمانت ترب",
                "وضعیت موجودی",
                "قیمت",
                "مقدار زمان تغییر قیمت",
                "واحد زمان تغییر قیمت",
                "لینک",
                "عکس",
            ]
            ws.append(headers)

            header_fill = PatternFill(
                start_color="1F4E78", end_color="1F4E78", fill_type="solid"
            )
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            fill_white = PatternFill(
                start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
            )
            fill_cream = PatternFill(
                start_color="FFF9E6", end_color="FFF9E6", fill_type="solid"
            )

            json_export_list = []
            captured_screenshots = []  # لیست موقت برای نگهداری آدرس عکس‌های گرفته‌شده
            use_cream = False

            chart_module = (
                TorobChartModule(excel_filename) if TorobChartModule else None
            )

            with sync_playwright() as p:
                launch_args = [
                    "--disable-blink-features=AutomationControlled",
                    "--no-sandbox",
                ]
                try:
                    browser = p.chromium.launch(
                        headless=is_headless, args=launch_args, channel="msedge"
                    )
                except Exception:  # noqa: BLE001
                    browser = p.chromium.launch(headless=is_headless, args=launch_args)

                context = browser.new_context(
                    user_agent=(
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
                        " (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
                    ),
                    locale="fa-IR",
                )
                page = context.new_page()

                page.goto(
                    "https://torob.com/", wait_until="domcontentloaded", timeout=60000
                )
                time.sleep(2)
                self.handle_captcha_if_present(page)

                for idx, (_, row) in enumerate(sliced_df.iterrows()):
                    if self.stop_requested:
                        break

                    product_name = str(row.get("نام", "")).strip()

                    if not product_name or pd.isna(product_name):
                        continue

                    input_values = []
                    for col in input_col:
                        val = row.get(col, "-")
                        if pd.isna(val):
                            val = "-"
                        input_values.append(val)

                    current_fill = fill_cream if use_cream else fill_white
                    use_cream = not use_cream

                    row_display = (
                        "تکی"
                        if mode != "استخراج از فایل اکسل"
                        else f"سطر {start_row + idx}"
                    )
                    self.log(f"\n🔍 [{row_display}] جستجو: {product_name}")

                    while True:
                        if self.stop_requested:
                            break

                        time.sleep(random.uniform(2, 5))

                        search_api_url = (
                            "https://api.torob.com/v4/base-product/search/?q="
                            f"{urllib.parse.quote(product_name)}&page=0&size=24"
                        )
                        search_data = self.safe_fetch_json(page, search_api_url)

                        results = (
                            search_data.get("results", [])
                            if isinstance(search_data, dict)
                            else []
                        )

                        if results:
                            break

                        self.log(f"❌ یافت نشد: {product_name} — احتمال مسدودیت ترب")
                        self.wait_for_unblock()

                    if self.stop_requested:
                        break

                    target_results = results[:max_results]

                    for top_product in target_results:
                        if self.stop_requested:
                            break

                        prk = top_product.get("random_key") or top_product.get("prk")
                        image_url = top_product.get("image_url", "-")

                        if prk:
                            sellers_url = f"https://api.torob.com/v4/base-product/sellers/?prk={prk}"
                            sellers_data = self.safe_fetch_json(page, sellers_url)

                            sellers_list = []
                            if sellers_data and isinstance(sellers_data, list):
                                sellers_list = sellers_data
                            elif sellers_data and isinstance(sellers_data, dict):
                                sellers_list = sellers_data.get("results", [])

                            if not sellers_list:
                                details_url = f"https://api.torob.com/v4/base-product/details/?prk={prk}"
                                details_data = self.safe_fetch_json(page, details_url)

                                if details_data and isinstance(details_data, dict):
                                    sellers_list = details_data.get(
                                        "products_info", {}
                                    ).get("result", [])

                            for seller in sellers_list:
                                torob_title = seller.get("name1") or top_product.get(
                                    "name1", "-"
                                )
                                guarantee_info = seller.get("guarantee_info") or {}
                                torob_guarantee = (
                                    "دارد"
                                    if guarantee_info.get("status") == "enabled"
                                    else "ندارد"
                                )
                                shop_name = seller.get("shop_name") or "-"
                                score_info = seller.get("score_info") or {}
                                score_and_history = (
                                    score_info.get("score_text")
                                    or seller.get("shop_score_text")
                                    or "-"
                                )
                                price_change = (
                                    seller.get("last_price_change_date") or "-"
                                )

                                clean_cpt_data, cpt_unit = self.cpt_data_cleaner(
                                    price_change
                                )

                                price_text = (
                                    seller.get("price_text")
                                    or seller.get("price_string")
                                    or "-"
                                )

                                clean_price_data, product_status = (
                                    self.price_data_cleaner(price_text)
                                )

                                link = seller.get("page_url") or "-"

                                row_data = input_values + [
                                    torob_title,
                                    shop_name,
                                    score_and_history,
                                    torob_guarantee,
                                    product_status,
                                    clean_price_data,
                                    clean_cpt_data,
                                    cpt_unit,
                                    link,
                                    image_url,
                                ]
                                ws.append(row_data)

                                for cell in ws[ws.max_row]:
                                    cell.fill = current_fill

                                json_export_dict1 = {
                                    "torob_title": torob_title,
                                    "shop_name": shop_name,
                                    "score_and_history": score_and_history,
                                    "torob_guarantee": torob_guarantee,
                                    "availability": product_status,
                                    "price": clean_price_data,
                                    "price_change": clean_cpt_data,
                                    "change_price_time_unit": cpt_unit,
                                    "link": link,
                                    "image_url": image_url,
                                    "prk": prk,
                                    "raw_seller_data": seller,
                                }

                                json_export_dict2 = {}

                                for col, val in zip(input_col, input_values):
                                    json_export_dict2[col] = val

                                json_export_dict2 = {
                                    **json_export_dict2,
                                    **json_export_dict1,
                                }

                                json_export_list.append(json_export_dict2)

                    self.log(f"✅ موفق: {product_name}")

                    # ثبت اسکرین‌شات و ذخیره مسیر در رم، بدون درگیری با فایل اکسل
                    if chart_module and target_results:
                        second_result = (
                            target_results[1]
                            if len(target_results) > 1
                            else target_results[0]
                        )
                        selected_prk = second_result.get(
                            "random_key"
                        ) or second_result.get("prk")
                        result_tag = (
                            "نتیجه دوم (واقعی)"
                            if len(target_results) > 1
                            else "تنها نتیجه"
                        )

                        if selected_prk:
                            self.log(
                                f"📸 ثبت اسکرین‌شات نمودار ({result_tag}): {product_name}"
                            )
                            prod_url = f"https://torob.com/p/{selected_prk}/"

                            try:
                                page.goto(
                                    prod_url,
                                    wait_until="domcontentloaded",
                                    timeout=30000,
                                )
                                time.sleep(2)

                                self.handle_captcha_if_present(page)

                                chart_btn = page.locator(
                                    'button[aria-label="لیست تغییرات قیمت"],'
                                    ' button[class*="chartOpenButton"]'
                                ).first

                                if chart_btn.count() > 0:
                                    chart_btn.scroll_into_view_if_needed()
                                    time.sleep(0.5)
                                    chart_btn.click()
                                    time.sleep(2)

                                    # مسیر عکس هم به پوشه اصلی منتقل شد تا در روت درایوها ایجاد نشود
                                    img_filename = os.path.join(
                                        script_dir, f"temp_chart_{idx}.png"
                                    )

                                    modal_element = page.locator(
                                        'div[role="dialog"]'
                                    ).first
                                    if (
                                        modal_element.count() > 0
                                        and modal_element.is_visible()
                                    ):
                                        modal_element.screenshot(path=img_filename)
                                    else:
                                        page.screenshot(path=img_filename)

                                    # فقط نام محصول و مسیر عکس در لیست موقت ذخیره می‌شود
                                    captured_screenshots.append(
                                        {"name": product_name, "path": img_filename}
                                    )

                                    page.keyboard.press("Escape")
                                    time.sleep(0.8)
                                else:
                                    self.log(
                                        f"⚠️ دکمه نمودار قیمت برای {product_name} یافت نشد."
                                    )

                            except Exception as err:  # noqa: BLE001
                                self.log(f"❌ خطا در اسکرین‌شات {product_name}: {err}")

                            time.sleep(1.5)

                # ۱. ذخیره قطعی متن‌ها روی هارد دیسک
                wb.save(excel_filename)

                with open(json_filename, "w", encoding="utf-8") as f:
                    json.dump(json_export_list, f, ensure_ascii=False, indent=2)

                # ۲. رسم نمودار و تزریق امن عکس‌ها به اکسلِ ساخته‌شده
                if chart_module:
                    chart_module.generate_charts()
                    self.log("📊 نمودارهای تحلیلی کلی با موفقیت به اکسل اضافه شدند.")

                    if captured_screenshots:
                        self.log("📸 در حال درج اسکرین‌شات‌های گرفته‌شده در فایل اکسل...")
                        image_row = 2
                        for img_data in captured_screenshots:
                            chart_module.add_chart_image(
                                img_data["name"], img_data["path"], image_row
                            )
                            image_row += 15

                            # حذف عکس‌های موقت پس از اتمام جایگذاری
                            if os.path.exists(img_data["path"]):
                                os.remove(img_data["path"])

                        self.log("✅ تمامی اسکرین‌شات‌ها با موفقیت در اکسل درج شدند.")

                browser.close()

            self.log(f"\n💾 فایل اکسل: {excel_filename}\n💾 فایل JSON: {json_filename}")

            if self.stop_requested:
                messagebox.showwarning(
                    "توقف اضطراری",
                    f"عملیات متوقف شد.\nاطلاعات تا این لحظه در فایل‌های زیر ذخیره گردید.\n\nاکسل: {excel_filename}",
                )
            else:
                messagebox.showinfo(
                    "پایان عملیات",
                    f"استخراج کامل شد.\n\nاکسل: {excel_filename}\nجیسون: {json_filename}",
                )

        except Exception as e:  # noqa: BLE001
            self.log(f"\n❌ خطا: {e}")
            messagebox.showerror("خطا", f"خطایی رخ داد: {e}")

        finally:
            self.is_running = False
            self.btn_start.configure(state="normal")
            self.btn_stop.configure(state="disabled")
            self.seg_button.configure(state="normal")
            self.btn_select_file.configure(state="normal")


if __name__ == "__main__":
    app = TorobScraperApp()
    app.mainloop()

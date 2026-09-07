import os
import openpyxl
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.drawing.image import Image
import pandas as pd


class TorobChartModule:

  def __init__(self, excel_filename):
    self.excel_filename = excel_filename

  def generate_charts(self):
    """ساخت نمودارهای میله‌ای و دایره‌ای در شیت دوم"""
    try:
      wb = openpyxl.load_workbook(self.excel_filename)

      if 'تحلیل و نمودارها' in wb.sheetnames:
        ws_chart = wb['تحلیل و نمودارها']
      else:
        ws_chart = wb.create_sheet(title='تحلیل و نمودارها')

      df = pd.read_excel(self.excel_filename, sheet_name='گزارش ترب')

      # ۱. نمودار ۱۰ فروشگاه برتر
      df_valid_shops = df[
          (df['فروشگاه'] != '-') & (df['فروشگاه'].notna())
      ].copy()
      shop_counts = (
          df_valid_shops['فروشگاه'].value_counts().head(10).reset_index()
      )
      shop_counts.columns = ['فروشگاه', 'تعداد محصول']

      ws_chart.cell(row=1, column=1, value='فروشگاه')
      ws_chart.cell(row=1, column=2, value='تعداد محصول')

      for idx, row in shop_counts.iterrows():
        ws_chart.cell(row=idx + 2, column=1, value=row['فروشگاه'])
        ws_chart.cell(row=idx + 2, column=2, value=row['تعداد محصول'])

      chart_bar = BarChart()
      chart_bar.type = 'col'
      chart_bar.style = 10
      chart_bar.title = '۱۰ فروشگاه برتر'
      chart_bar.y_axis.title = 'تعداد'

      data_ref = Reference(
          ws_chart, min_col=2, min_row=1, max_row=len(shop_counts) + 1
      )
      cats_ref = Reference(
          ws_chart, min_col=1, min_row=2, max_row=len(shop_counts) + 1
      )

      chart_bar.add_data(data_ref, titles_from_data=True)
      chart_bar.set_categories(cats_ref)
      chart_bar.width = 16
      chart_bar.height = 9
      ws_chart.add_chart(chart_bar, 'D2')

      # ۲. نمودار وضعیت ضمانت ترب
      guarantee_counts = df['ضمانت ترب'].value_counts().reset_index()
      guarantee_counts.columns = ['وضعیت', 'تعداد']

      start_row_g = len(shop_counts) + 4
      ws_chart.cell(row=start_row_g, column=1, value='وضعیت ضمانت')
      ws_chart.cell(row=start_row_g, column=2, value='تعداد')

      for idx, row in guarantee_counts.iterrows():
        ws_chart.cell(row=start_row_g + 1 + idx, column=1, value=row['وضعیت'])
        ws_chart.cell(row=start_row_g + 1 + idx, column=2, value=row['تعداد'])

      chart_pie = PieChart()
      chart_pie.title = 'وضعیت ضمانت ترب'

      data_pie = Reference(
          ws_chart,
          min_col=2,
          min_row=start_row_g,
          max_row=start_row_g + len(guarantee_counts),
      )
      cats_pie = Reference(
          ws_chart,
          min_col=1,
          min_row=start_row_g + 1,
          max_row=start_row_g + len(guarantee_counts),
      )

      chart_pie.add_data(data_pie, titles_from_data=True)
      chart_pie.set_categories(cats_pie)
      chart_pie.width = 12
      chart_pie.height = 8
      ws_chart.add_chart(chart_pie, 'D17')

      wb.save(self.excel_filename)
      return True
    except Exception as e:
      print(f'خطا در ساخت نمودارها: {e}')
      return False

  def add_chart_image(self, product_name, image_path, row_index):
    """درج اسکرین‌شات نمودار تغییر قیمت در شیت دوم"""
    try:
      wb = openpyxl.load_workbook(self.excel_filename)
      ws_chart = (
          wb['تحلیل و نمودارها']
          if 'تحلیل و نمودارها' in wb.sheetnames
          else wb.create_sheet(title='تحلیل و نمودارها')
      )

      start_col = 12  # ستون L در اکسل
      ws_chart.cell(
          row=row_index,
          column=start_col,
          value=f'📈 نمودار تغییر قیمت: {product_name[:35]}',
      )

      img = Image(image_path)
      img.width = 450
      img.height = 250

      cell_pos = f'L{row_index + 1}'
      ws_chart.add_image(img, cell_pos)

      wb.save(self.excel_filename)
      return True
    except Exception as e:
      print(f'خطا در افزودن تصویر نمودار: {e}')
      return False